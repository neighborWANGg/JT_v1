import hashlib
import math
import re
from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path
from typing import Protocol
from zipfile import BadZipFile

import httpx
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from .config import settings
from .models import Message, SourceFile

REQUIRED_COLUMNS = ("序号", "时间", "群ID", "微信号", "聊天内容", "IP登录地址")
COUNTER_TERMS = ("没", "未", "不是", "取消", "别动", "等通知")
FIELD_PATTERN = re.compile(r"wxid_[A-Za-z0-9]+|(?:\d{1,3}\.){3}\d{1,3}|(?:群ID|群号)[：:\s]*([\w-]+)", re.I)


class Embedder(Protocol):
    def encode(self, texts: list[str]) -> list[list[float]]: ...


class HashEmbedder:
    """供开发和测试使用，不代表语义相似度。"""

    def encode(self, texts: list[str]) -> list[list[float]]:
        return [self._encode_one(text) for text in texts]

    @staticmethod
    def _encode_one(text: str) -> list[float]:
        raw = bytearray()
        counter = 0
        while len(raw) < settings.embedding_dimension:
            raw.extend(hashlib.sha256(f"{counter}:{text}".encode("utf-8")).digest())
            counter += 1
        vector = [(value / 127.5) - 1 for value in raw[: settings.embedding_dimension]]
        norm = math.sqrt(sum(value * value for value in vector)) or 1
        return [value / norm for value in vector]


class BgeEmbedder:
    def __init__(self) -> None:
        from FlagEmbedding import BGEM3FlagModel

        self.model = BGEM3FlagModel(settings.bge_model_path, use_fp16=True)

    def encode(self, texts: list[str]) -> list[list[float]]:
        output = self.model.encode(
            texts,
            batch_size=8,
            max_length=512,
            return_dense=True,
            return_sparse=False,
            return_colbert_vecs=False,
        )
        return output["dense_vecs"].tolist()


class LmStudioEmbedder:
    def __init__(self, transport: httpx.BaseTransport | None = None) -> None:
        self.client = httpx.Client(timeout=60, transport=transport)

    def encode(self, texts: list[str]) -> list[list[float]]:
        response = self.client.post(
            f"{settings.embedding_base_url}/embeddings",
            json={"model": settings.embedding_model, "input": texts},
        )
        response.raise_for_status()
        return [
            item["embedding"]
            for item in sorted(response.json()["data"], key=lambda item: item["index"])
        ]


def build_embedder() -> Embedder:
    if settings.embedding_backend == "bge":
        return BgeEmbedder()
    if settings.embedding_backend == "lmstudio":
        return LmStudioEmbedder()
    return HashEmbedder()


embedder: Embedder = build_embedder()


def stable_message_id(file_hash: str, sheet: str, row_number: int) -> str:
    return hashlib.sha256(f"{file_hash}:{sheet}:{row_number}".encode("utf-8")).hexdigest()


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _datetime(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = _text(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError as exc:
        raise ValueError(f"无法解析时间：{text}") from exc


def import_xlsx(session: Session, case_id, filename: str, payload: bytes) -> tuple[SourceFile, int, bool]:
    file_hash = hashlib.sha256(payload).hexdigest()
    existing = session.scalar(
        select(SourceFile).where(SourceFile.case_id == case_id, SourceFile.sha256 == file_hash)
    )
    if existing:
        count = session.query(Message).filter(Message.source_file_id == existing.id).count()
        return existing, count, True

    try:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException) as exc:
        raise ValueError("文件不是有效的 xlsx 工作簿") from exc
    if "聊天记录" not in workbook.sheetnames:
        raise ValueError("Excel 缺少工作表：聊天记录")
    sheet = workbook["聊天记录"]
    rows = sheet.iter_rows(values_only=True)
    headers = list(next(rows, []))
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise ValueError(f"Excel 缺少字段：{', '.join(missing)}")
    column = {name: headers.index(name) for name in REQUIRED_COLUMNS}

    case_dir = settings.data_dir / str(case_id)
    case_dir.mkdir(parents=True, exist_ok=True)
    suffix = Path(filename).suffix.lower() or ".xlsx"
    stored = case_dir / f"{file_hash}{suffix}"
    stored.write_bytes(payload)
    stored.chmod(0o444)

    source = SourceFile(
        case_id=case_id,
        filename=Path(filename).name,
        sha256=file_hash,
        stored_path=str(stored),
    )
    session.add(source)
    session.flush()

    parsed = []
    for row_number, row in enumerate(rows, start=2):
        content = _text(row[column["聊天内容"]])
        if content:
            parsed.append((row_number, row, content))
    vectors = embedder.encode([item[2] for item in parsed])
    if len(vectors) != len(parsed):
        raise RuntimeError("向量数量与消息数量不一致")

    for (row_number, row, content), vector in zip(parsed, vectors, strict=True):
        session.add(Message(
            id=stable_message_id(file_hash, sheet.title, row_number),
            case_id=case_id,
            source_file_id=source.id,
            source_row=row_number,
            sequence=_text(row[column["序号"]]),
            sent_at=_datetime(row[column["时间"]]),
            group_id=_text(row[column["群ID"]]),
            sender_id=_text(row[column["微信号"]]),
            content=content,
            ip_address=_text(row[column["IP登录地址"]]),
            embedding=vector,
        ))
    session.commit()
    return source, len(parsed), False


def _terms(text: str) -> set[str]:
    normalized = re.sub(r"[^\w\u4e00-\u9fff]+", "", text.lower())
    ascii_terms = set(re.findall(r"[a-z0-9_]{2,}", normalized))
    chinese = "".join(re.findall(r"[\u4e00-\u9fff]", normalized))
    return ascii_terms | {chinese[index:index + 2] for index in range(max(0, len(chinese) - 1))}


def search_messages(
    session: Session,
    case_id,
    question: str,
    limit: int | None = None,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
    sender_id: str | None = None,
    group_id: str | None = None,
    ip_address: str | None = None,
    source_file_id=None,
):
    query_vector = embedder.encode([question])[0]
    statement = select(Message).where(Message.case_id == case_id)
    if start_date:
        statement = statement.where(Message.sent_at >= datetime.combine(start_date, time.min))
    if end_date:
        statement = statement.where(Message.sent_at <= datetime.combine(end_date, time.max))
    for field, value in (
        (Message.sender_id, sender_id),
        (Message.group_id, group_id),
        (Message.ip_address, ip_address),
        (Message.source_file_id, source_file_id),
    ):
        if value is not None:
            statement = statement.where(field == value)
    messages = session.scalars(statement).all()
    query_terms = _terms(question)
    field_values = {
        (match.group(1) or match.group(0)).strip().lower()
        for match in FIELD_PATTERN.finditer(question)
    }

    # ponytail: SQLite 原型在应用内扫描；超过约 10 万条消息时迁移 pgvector/OpenSearch。
    scored = []
    for message in messages:
        similarity = sum(a * b for a, b in zip(query_vector, message.embedding, strict=True))
        message_terms = _terms(message.content)
        lexical = len(query_terms & message_terms) / max(1, len(query_terms))
        fields = {value.lower() for value in (message.sender_id, message.group_id, message.ip_address) if value}
        field_bonus = 2.0 * len(field_values & fields)
        counter_bonus = 0.6 * sum(term in question and term in message.content for term in COUNTER_TERMS)
        exact_bonus = 0.5 if question.strip().lower() in message.content.lower() else 0.0
        scored.append((message, similarity + lexical + field_bonus + counter_bonus + exact_bonus))
    ranked = sorted(scored, key=lambda item: item[1], reverse=True)
    if limit is None:
        return ranked
    return ranked[:limit]


def matches_filters(
    message: Message,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
    sender_id: str | None = None,
    group_id: str | None = None,
    ip_address: str | None = None,
    source_file_id=None,
) -> bool:
    return not any((
        start_date and (not message.sent_at or message.sent_at.date() < start_date),
        end_date and (not message.sent_at or message.sent_at.date() > end_date),
        sender_id is not None and message.sender_id != sender_id,
        group_id is not None and message.group_id != group_id,
        ip_address is not None and message.ip_address != ip_address,
        source_file_id is not None and message.source_file_id != source_file_id,
    ))


def context_window(
    session: Session,
    case_id,
    message_id: str,
    before: int,
    after: int,
) -> list[tuple[Message, int]]:
    anchor = session.get(Message, message_id)
    if not anchor or anchor.case_id != case_id:
        raise LookupError("消息不存在")

    statement = select(Message).where(
        Message.case_id == case_id,
        Message.source_file_id == anchor.source_file_id,
    )
    statement = statement.where(
        Message.group_id == anchor.group_id
        if anchor.group_id is not None
        else Message.group_id.is_(None)
    )
    messages = session.scalars(statement.order_by(Message.source_row)).all()
    anchor_index = next(index for index, message in enumerate(messages) if message.id == message_id)
    start = max(0, anchor_index - before)
    end = min(len(messages), anchor_index + after + 1)
    return [(messages[index], index - anchor_index) for index in range(start, end)]


def time_neighbors(
    session: Session,
    case_id,
    message_id: str,
    minutes: int,
    limit: int,
) -> list[tuple[Message, int]]:
    anchor = session.get(Message, message_id)
    if not anchor or anchor.case_id != case_id:
        raise LookupError("消息不存在")
    if anchor.sent_at is None:
        return [(anchor, 0)]

    statement = select(Message).where(
        Message.case_id == case_id,
        Message.sent_at.between(
            anchor.sent_at - timedelta(minutes=minutes),
            anchor.sent_at + timedelta(minutes=minutes),
        ),
    )
    statement = statement.where(
        Message.group_id == anchor.group_id
        if anchor.group_id is not None
        else Message.source_file_id == anchor.source_file_id
    )
    messages = session.scalars(
        statement.order_by(Message.sent_at, Message.source_row).limit(limit)
    ).all()
    return [
        (message, round((message.sent_at - anchor.sent_at).total_seconds()))
        for message in messages
    ]


def expand_context(
    session: Session,
    case_id,
    scored_messages: list[tuple[Message, float]],
    before: int,
    after: int,
) -> list[tuple[Message, float]]:
    expanded: dict[str, tuple[Message, float]] = {
        message.id: (message, score) for message, score in scored_messages
    }
    for anchor, anchor_score in scored_messages:
        for message, distance in context_window(session, case_id, anchor.id, before, after):
            context_score = anchor_score - abs(distance) * 0.01
            current = expanded.get(message.id)
            if current is None or context_score > current[1]:
                expanded[message.id] = (message, context_score)
    return sorted(
        expanded.values(),
        key=lambda item: (
            item[0].sent_at or datetime.min,
            item[0].source_row,
        ),
    )
