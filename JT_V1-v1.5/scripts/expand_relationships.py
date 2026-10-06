from pathlib import Path
import sys

from openpyxl import load_workbook


RELATIONS = {
    "周凯": ["刘桂香", "刘洋", "周建国", "林芳", "周琳", "吴磊", "孙浩", "郑浩", "钱多多", "王雅", "赵立国", "陈晓东", "方敏", "蒋涛", "韩雪", "何平"],
    "刘桂香": ["周凯", "周建国", "林芳", "周琳", "刘洋", "王秀兰", "陈淑芬", "李梅", "张阿姨", "赵叔叔", "孙阿姨", "邻居老许", "社区小林", "药房小周", "张师傅", "李医生"],
    "刘洋": ["周凯", "刘桂香", "赵立国", "王雅", "陈晓东", "方敏", "蒋涛", "韩雪", "何平", "孙浩", "吴磊", "郑浩", "钱多多", "彭程", "高远", "林工"],
}


def expand(path: Path) -> None:
    workbook = load_workbook(path)
    if "人员关系" in workbook.sheetnames:
        del workbook["人员关系"]
    sheet = workbook.create_sheet("人员关系")
    sheet.append(["序号", "主体", "关系人", "关系类型", "微信号", "IP地址", "关系说明"])
    index = 1
    for subject, names in RELATIONS.items():
        for relation_index, name in enumerate(names, 1):
            wechat = f"wxid_{subject.encode('punycode').decode()[:4].lower()}_{relation_index:02d}"
            sheet.append([index, subject, name, "亲属" if relation_index <= 5 else "同事/朋友", wechat, f"10.24.{(index % 200) + 1}.{(index % 240) + 10}", f"{subject}的第{relation_index}位关系人"])
            index += 1
    sheet.freeze_panes = "A2"
    workbook.save(path)


if __name__ == "__main__":
    for value in sys.argv[1:]:
        target = Path(value)
        if target.is_dir():
            for workbook_path in target.glob("*.xlsx"):
                if workbook_path.name.startswith("~$"):
                    continue
                expand(workbook_path)
        else:
            expand(target)
