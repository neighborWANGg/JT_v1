"""人员总结、关键词检索、群聊研判算法。"""

from __future__ import annotations

import math
import re
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .config import settings
from .mysql_repo import MysqlChat, MysqlPerson
from .schemas import SearchHit

SELF_REMARKS = ("", "本人", "主体", "主要人员")
QUESTION_WORDS = ("有没有", "是否", "有无", "请问", "帮我", "查询", "查找", "检索", "分析", "判断", "说明", "消息中", "聊天中", "数据中", "记录中", "吗", "么", "呢")
PROFILE_QUESTION = re.compile(
    r"总结|介绍|概况|档案|是谁|什么人|人员信息|个人信息|关系人|家庭关系|画像|群类型|什么群|哪些群|群研判"
)


@dataclass
class ChatItem:
    message_id: str
    source_file_id: uuid.UUID
    source_row: int
    sent_at: datetime | None
    group_id: str | None
    sender_id: str | None
    content: str
    ip_address: str | None
    embedding: list[float] = field(default_factory=list)


@dataclass
class GroupProfile:
    group_id: str
    member_count: int
    members: list[str]
    message_count: int
    subject_speak_count: int
    related_count: int
    related_names: list[str]
    messages: list[dict] = field(default_factory=list)
    category: str = "待研判"
    reason: str = ""


def resolve_bge_path() -> Path | None:
    candidates = [
        Path(settings.bge_model_path),
        Path(__file__).resolve().parents[3] / "model" / "bge-m3",
        Path(__file__).resolve().parents[1] / "model" / "bge-m3",
    ]
    for path in candidates:
        if path.exists() and (path / "config.json").exists():
            return path
    return None


_bge_model = None


def encode_bge(texts: list[str]) -> list[list[float]]:
    global _bge_model
    path = resolve_bge_path()
    if path is None:
        raise FileNotFoundError("未找到本地 BGE-M3 模型")
    if _bge_model is None:
        from FlagEmbedding import BGEM3FlagModel

        kwargs = {"use_fp16": True, "query_max_length": 512}
        try:
            import torch

            if torch.cuda.is_available():
                kwargs["devices"] = "cuda:0"
        except ImportError:
            pass
        _bge_model = BGEM3FlagModel(str(path), **kwargs)
    output = _bge_model.encode(
        texts,
        batch_size=8,
        max_length=512,
        return_dense=True,
        return_sparse=False,
        return_colbert_vecs=False,
    )
    return output["dense_vecs"].tolist()


def cosine(left: list[float], right: list[float]) -> float:
    if not left or not right or len(left) != len(right):
        return 0.0
    score = sum(a * b for a, b in zip(left, right, strict=True))
    norm_l = math.sqrt(sum(a * a for a in left)) or 1
    norm_r = math.sqrt(sum(b * b for b in right)) or 1
    return score / (norm_l * norm_r)


def is_subject(person: MysqlPerson) -> bool:
    remark = (person.remark or "").strip()
    return remark in SELF_REMARKS


def subjects_of(persons: list[MysqlPerson]) -> list[MysqlPerson]:
    found = [person for person in persons if is_subject(person)]
    return found or persons[:1]


def related_of(persons: list[MysqlPerson], subjects: list[MysqlPerson]) -> list[MysqlPerson]:
    subject_ids = {person.id_card for person in subjects}
    return [person for person in persons if person.id_card not in subject_ids]


def relationship_rows_from_sources(paths: list[Path]) -> list[dict]:
    """从案件源文件中提取规范化人员关系，当前兼容两种既有 Excel 表头。"""
    result: list[dict] = []
    seen: set[tuple[str, str, str, str]] = set()
    for path in paths:
        if path.suffix.lower() not in {".xlsx", ".xlsm"} or not path.exists():
            continue
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception:
            continue
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows, [])]
            columns = {name: index for index, name in enumerate(headers) if name}

            def value(row, *names: str) -> str:
                for name in names:
                    index = columns.get(name)
                    if index is not None and index < len(row) and row[index] is not None:
                        return str(row[index]).strip()
                return ""

            expanded = any(name in columns for name in ("主体", "当事人")) and any(
                name in columns for name in ("关系人", "关联人")
            )
            if expanded:
                for row in rows:
                    subject = value(row, "主体", "当事人")
                    related = value(row, "关系人", "关联人")
                    relation = value(row, "关系类型", "关系") or "关系人"
                    if not subject or not related:
                        continue
                    item = {
                        "subject": subject,
                        "related": related,
                        "relation": relation,
                        "wechat_id": value(row, "微信号", "微信ID", "账号"),
                        "ip_address": value(row, "IP地址", "IP登录地址"),
                        "description": value(row, "关系说明", "备注"),
                    }
                    key = (subject, related, relation, item["wechat_id"])
                    if key not in seen:
                        result.append(item)
                        seen.add(key)
                continue

            if "姓名" not in columns or not any(name in columns for name in ("关系", "关系类型")):
                continue
            simple_rows = []
            for row in rows:
                name = value(row, "姓名")
                if name:
                    simple_rows.append((row, name, value(row, "关系", "关系类型")))
            subject = next((name for _, name, relation in simple_rows if relation in SELF_REMARKS), "")
            subject = subject or (simple_rows[0][1] if simple_rows else "")
            for row, name, relation in simple_rows:
                if not subject or name == subject or relation in SELF_REMARKS:
                    continue
                item = {
                    "subject": subject,
                    "related": name,
                    "relation": relation or "关系人",
                    "wechat_id": value(row, "微信号", "微信ID", "账号"),
                    "ip_address": value(row, "IP地址", "IP登录地址"),
                    "description": value(row, "关系说明", "备注"),
                }
                key = (subject, name, item["relation"], item["wechat_id"])
                if key not in seen:
                    result.append(item)
                    seen.add(key)
        workbook.close()
    return result


def apply_manual_relations(imported: list[dict], stored: list[dict]) -> tuple[list[dict], set[tuple[str, str]], dict[tuple[str, str], str], set[str]]:
    rows = [
        {
            "subject": str(row.get("subject") or "").strip(),
            "related": str(row.get("related") or "").strip(),
            "relation": str(row.get("relation") or "关系人").strip() or "关系人",
            "wechat_id": str(row.get("wechat_id") or "").strip(),
            "ip_address": row.get("ip_address"),
            "description": row.get("description"),
        }
        for row in imported
        if str(row.get("subject") or "").strip() and str(row.get("related") or "").strip()
    ]
    hidden: set[tuple[str, str]] = set()
    hidden_names: set[str] = set()
    overrides: dict[tuple[str, str], str] = {}
    for item in stored:
        source = str(item.get("source_name") or "").strip()
        target = str(item.get("target_name") or "").strip()
        relation = str(item.get("relation") or "关系人").strip() or "关系人"
        if not source or not target:
            continue
        pair = (source, target)
        reverse = (target, source)
        if item.get("removed") and source == target:
            hidden_names.add(source)
            continue
        if item.get("removed"):
            hidden.add(pair)
            hidden.add(reverse)
            continue
        rows = [row for row in rows if (row["subject"], row["related"]) not in {pair, reverse}]
        rows.append({
            "subject": source,
            "related": target,
            "relation": relation,
            "wechat_id": str(item.get("wechat_id") or "").strip(),
            "ip_address": item.get("ip_address"),
            "description": item.get("description"),
        })
        overrides[pair] = relation
        overrides[reverse] = relation
    rows = [
        row for row in rows
        if (row["subject"], row["related"]) not in hidden
        and row["subject"] not in hidden_names
        and row["related"] not in hidden_names
    ]
    return rows, hidden, overrides, hidden_names


def wechat_keys(*values) -> set[str]:
    return {str(value).strip().lower() for value in values if str(value or "").strip()}


def person_wechat_keys(person: MysqlPerson) -> set[str]:
    keys = wechat_keys(person.wechat_id, person.wechat_alias, person.original_wechat_id)
    for account in person.accounts or ():
        keys |= wechat_keys(account.wechat_id, account.wechat_alias, account.original_wechat_id)
    return keys


UPLOADED_REMARKS = ("本人", "主体", "主要人员")


def uploaded_people_index(persons: list[MysqlPerson], friend_rows: list[dict] | None = None) -> dict[str, str]:
    """已上传数据里能确定身份的人：文件归属人和标记为本人的人员，不含好友表里的联系人。"""
    index: dict[str, str] = {}
    for person in persons:
        remark = (person.remark or "").strip()
        name = (person.name or "").strip()
        if remark not in UPLOADED_REMARKS or not name:
            continue
        for key in person_wechat_keys(person):
            index.setdefault(key, name)
    for row in friend_rows or []:
        name = str(row.get("owner_name") or "").strip()
        if not name:
            continue
        for key in wechat_keys(row.get("owner_wechat_id"), row.get("owner_alias")):
            index.setdefault(key, name)
    return index


def groups_with_uploaded_people(messages: list[ChatItem], index: dict[str, str]) -> dict[str, list[str]]:
    found: dict[str, set[str]] = defaultdict(set)
    for item in messages:
        sender = str(item.sender_id or "").strip().lower()
        name = index.get(sender)
        if item.group_id and name:
            found[item.group_id].add(name)
    return {group_id: sorted(names) for group_id, names in found.items() if len(names) >= 2}


def account_payload(person: MysqlPerson) -> list[dict]:
    accounts = [
        {
            "wechat_id": account.wechat_id,
            "wechat_alias": account.wechat_alias,
            "original_wechat_id": account.original_wechat_id,
            "nickname": account.nickname or person.wx_nickname,
        }
        for account in person.accounts or ()
        if account.wechat_id or account.wechat_alias or account.original_wechat_id
    ]
    if accounts:
        return accounts
    if person.wechat_id or person.wechat_alias or person.original_wechat_id:
        return [{
            "wechat_id": person.wechat_id,
            "wechat_alias": person.wechat_alias,
            "original_wechat_id": person.original_wechat_id,
            "nickname": person.wx_nickname,
        }]
    return []


WECHAT_SLOTS = ("id", "alias", "original")
FRIEND_WECHAT_FIELDS = {
    "id": "friend_wechat_id",
    "alias": "friend_alias",
    "original": "friend_original_wechat_id",
}


def _wechat_token(value) -> str:
    return str(value or "").strip().lower()


def friend_identity(row: dict) -> dict:
    return {slot: _wechat_token(row.get(field)) for slot, field in FRIEND_WECHAT_FIELDS.items()}


def identities_match(left: dict, right: dict) -> bool:
    """微信ID对微信ID，微信号对微信号，原始账号对原始账号。某一侧该栏为空就跳过这一栏。"""
    if not any(left.values()) or not any(right.values()):
        return False
    for slot in WECHAT_SLOTS:
        if left[slot] and right[slot] and left[slot] == right[slot]:
            return True
    return False


def person_identities(person: MysqlPerson) -> list[dict]:
    identities = []
    if person.wechat_id or person.wechat_alias or person.original_wechat_id:
        identities.append(friend_identity({
            "friend_wechat_id": person.wechat_id,
            "friend_alias": person.wechat_alias,
            "friend_original_wechat_id": person.original_wechat_id,
        }))
    for account in person.accounts or ():
        if account.wechat_id or account.wechat_alias or account.original_wechat_id:
            identities.append(friend_identity({
                "friend_wechat_id": account.wechat_id,
                "friend_alias": account.wechat_alias,
                "friend_original_wechat_id": account.original_wechat_id,
            }))
    return identities


def friend_cluster(row: dict) -> dict:
    identity = friend_identity(row)
    return {"identity": identity, "row": row}


def _distinct_wechat(rows: list[dict], field: str) -> str | None:
    values = []
    seen = set()
    for row in rows:
        text = str(row.get(field) or "").strip()
        if text and text.lower() not in seen:
            seen.add(text.lower())
            values.append(text)
    return "、".join(values) or None


def _bucket_matches(identity: dict, bucket: dict) -> bool:
    return any(identities_match(identity, item) for item in bucket["identities"])


def owner_matches_person(row: dict, person: MysqlPerson) -> bool:
    owner_keys = wechat_keys(row.get("owner_wechat_id"), row.get("owner_alias"))
    if owner_keys and owner_keys & person_wechat_keys(person):
        return True
    owner_name = str(row.get("owner_name") or "").strip()
    return bool(owner_name and owner_name == (person.name or "").strip())


def friends_of(person: MysqlPerson, friend_rows: list[dict]) -> list[dict]:
    return [row for row in friend_rows if owner_matches_person(row, person)]


def person_in_friend_list(person: MysqlPerson, clusters: list[dict]) -> bool:
    identities = person_identities(person)
    return any(_bucket_matches(cluster["identity"], {"identities": identities}) for cluster in clusters)


def screen_wechat_friends(persons: list[MysqlPerson], friend_rows: list[dict], names: list[str]) -> dict:
    selected = []
    seen = set()
    for name in names:
        name = str(name or "").strip()
        matches = [person for person in persons if person.name == name]
        if not matches:
            continue
        for person in matches:
            marker = (person.id_card, person.name)
            if marker in seen:
                continue
            seen.add(marker)
            selected.append(person)
    pair_results = []
    for index, left in enumerate(selected):
        left_clusters = [friend_cluster(row) for row in friends_of(left, friend_rows)]
        for right in selected[index + 1:]:
            right_clusters = [friend_cluster(row) for row in friends_of(right, friend_rows)]
            left_has_right = person_in_friend_list(right, left_clusters)
            right_has_left = person_in_friend_list(left, right_clusters)
            if left_has_right and right_has_left:
                relation = "互为微信好友"
            elif left_has_right:
                relation = f"{left.name}的好友包含{right.name}"
            elif right_has_left:
                relation = f"{right.name}的好友包含{left.name}"
            else:
                relation = "未发现好友关系"
            pair_results.append({
                "left": left.name,
                "right": right.name,
                "left_has_right": left_has_right,
                "right_has_left": right_has_left,
                "mutual": left_has_right and right_has_left,
                "relation": relation,
            })
    common = []
    if len(selected) >= 2:
        own_identities = [identity for person in selected for identity in person_identities(person)]
        buckets: list[dict] = []
        for person in selected:
            for row in friends_of(person, friend_rows):
                cluster = friend_cluster(row)
                if not any(cluster["identity"].values()):
                    continue
                if _bucket_matches(cluster["identity"], {"identities": own_identities}):
                    continue
                matched = next((bucket for bucket in buckets if _bucket_matches(cluster["identity"], bucket)), None)
                if matched is None:
                    buckets.append({
                        "identities": [cluster["identity"]],
                        "rows": [dict(cluster["row"])],
                        "owners": {person.name},
                    })
                    continue
                matched["identities"].append(cluster["identity"])
                matched["rows"].append(dict(cluster["row"]))
                matched["owners"].add(person.name)
        seen_keys: set[str] = set()
        for bucket in buckets:
            if len(bucket["owners"]) < 2:
                continue
            marker = ",".join(sorted(
                "|".join(item[slot] for slot in WECHAT_SLOTS) for item in bucket["identities"]
            ))
            if not marker or marker in seen_keys:
                continue
            seen_keys.add(marker)
            rows = bucket["rows"]
            common.append({
                "nickname": next((row.get("friend_nickname") for row in rows if row.get("friend_nickname")), None),
                "remark": next((row.get("friend_remark") for row in rows if row.get("friend_remark")), None),
                "wechat_id": _distinct_wechat(rows, "friend_wechat_id"),
                "wechat_alias": _distinct_wechat(rows, "friend_alias"),
                "original_wechat_id": _distinct_wechat(rows, "friend_original_wechat_id"),
                "shared_count": len(bucket["owners"]),
                "shared_by": sorted(bucket["owners"]),
            })
        common.sort(key=lambda item: (-item["shared_count"], item.get("nickname") or "", item.get("wechat_id") or ""))
    return {
        "people": [person.name for person in selected],
        "missing": [name for name in names if str(name or "").strip() and str(name).strip() not in {person.name for person in selected}],
        "pairs": pair_results,
        "common_friends": common,
        "owner_counts": [
            {
                "name": person.name,
                "account_count": max(1, len(account_payload(person))),
                "friend_count": len(friends_of(person, friend_rows)),
                "wechat_ids": sorted(person_wechat_keys(person)),
            }
            for person in selected
        ],
    }


def build_case_graph(
    persons: list[MysqlPerson],
    items: list[ChatItem],
    relationship_rows: list[dict] | None = None,
    friend_rows: list[dict] | None = None,
    hidden_pairs: set[tuple[str, str]] | None = None,
    relation_overrides: dict[tuple[str, str], str] | None = None,
    hidden_names: set[str] | None = None,
) -> dict:
    """把人员库、关系表和群聊消息聚合为前端可直接渲染的案件图谱。"""
    relationship_rows = relationship_rows or []
    friend_rows = friend_rows or []
    hidden_pairs = hidden_pairs or set()
    relation_overrides = relation_overrides or {}
    hidden_names = hidden_names or set()
    nodes: dict[str, dict] = {}
    name_to_id: dict[str, str] = {}
    wechat_to_id: dict[str, str] = {}
    edges: list[dict] = []
    edge_keys: set[tuple[str, str, str]] = set()
    subject_names = {str(row.get("subject") or "").strip() for row in relationship_rows}
    subject_names.discard("")
    subject_names.update(person.name for person in subjects_of(persons) if person.name)

    def apply_profile(node: dict, details: dict) -> None:
        for key in (
            "gender", "id_card", "delivery_address", "takeout_address", "wechat_alias",
            "original_wechat_id", "wx_nickname", "phone", "ip_address", "description",
        ):
            if details.get(key) and not node.get(key):
                node[key] = details[key]
        if details.get("accounts"):
            existing = {(item.get("wechat_id"), item.get("wechat_alias"), item.get("original_wechat_id")) for item in node.get("accounts") or []}
            merged = list(node.get("accounts") or [])
            for account in details["accounts"]:
                marker = (account.get("wechat_id"), account.get("wechat_alias"), account.get("original_wechat_id"))
                if marker not in existing:
                    merged.append(account)
                    existing.add(marker)
            node["accounts"] = merged

    def ensure_node(name: str, wechat_id: str = "", role: str = "related", **details) -> str:
        name = (name or wechat_id or "未知账号").strip()
        wechat_id = (wechat_id or "").strip()
        node_id = wechat_to_id.get(wechat_id) if wechat_id else None
        node_id = node_id or name_to_id.get(name) or (f"wx:{wechat_id}" if wechat_id else f"name:{name}")
        node = nodes.setdefault(node_id, {
            "id": node_id,
            "name": name,
            "wechat_id": wechat_id or None,
            "role": role,
            "message_count": 0,
            "relation_count": 0,
            "ip_address": details.get("ip_address") or None,
            "description": details.get("description") or None,
            "gender": None,
            "id_card": None,
            "delivery_address": None,
            "takeout_address": None,
            "wechat_alias": None,
            "original_wechat_id": None,
            "wx_nickname": None,
            "phone": None,
            "accounts": [],
            "person_id": details.get("person_id") or None,
        })
        if details.get("person_id"):
            node["person_id"] = details["person_id"]
        if role == "subject" or node["role"] != "subject" and role == "participant":
            node["role"] = role
        if wechat_id and not node["wechat_id"]:
            node["wechat_id"] = wechat_id
        apply_profile(node, details)
        name_to_id[name] = node_id
        if wechat_id:
            wechat_to_id[wechat_id] = node_id
        for account in details.get("accounts") or []:
            for value in (account.get("wechat_id"), account.get("wechat_alias"), account.get("original_wechat_id")):
                if value:
                    wechat_to_id[str(value).strip()] = node_id
        return node_id

    def add_edge(source_name: str, target_name: str, relation: str, target_wechat: str = "", **details) -> None:
        pair = ((source_name or "").strip(), (target_name or "").strip())
        if pair[0] in hidden_names or pair[1] in hidden_names:
            return
        if pair in hidden_pairs or (pair[1], pair[0]) in hidden_pairs:
            return
        source = ensure_node(source_name, role="subject" if source_name in subject_names else "related")
        target = ensure_node(
            target_name,
            target_wechat,
            role="subject" if target_name in subject_names else "related",
            **details,
        )
        key = (source, target, relation or "关系人")
        if source != target and key not in edge_keys:
            payload = {"source": source, "target": target, "relation": key[2]}
            if details.get("kind"):
                payload["kind"] = details["kind"]
                payload["mutual"] = bool(details.get("mutual"))
            edges.append(payload)
            edge_keys.add(key)

    for person in persons:
        if person.name in hidden_names:
            continue
        role = "subject" if person.name in subject_names or is_subject(person) else "related"
        ensure_node(
            person.name,
            person.wechat_id or "",
            role=role,
            ip_address=person.ip_address,
            description=person.remark,
            gender=person.gender,
            id_card=person.id_card,
            delivery_address=person.delivery_address,
            takeout_address=person.takeout_address,
            wechat_alias=person.wechat_alias,
            original_wechat_id=person.original_wechat_id,
            wx_nickname=person.wx_nickname,
            phone=person.phone,
            accounts=account_payload(person),
            person_id=person.person_id,
        )
    fallback_subject = next(iter(subject_names), persons[0].name if persons else "")
    for person in related_of(persons, subjects_of(persons)):
        match = re.match(r"^(.+?)的(.+)$", (person.remark or "").strip())
        owner = match.group(1) if match else fallback_subject
        relation = relation_overrides.get((owner, person.name)) or (match.group(2) if match else (person.remark or "关系人"))
        if owner:
            add_edge(owner, person.name, relation, person.wechat_id or "", ip_address=person.ip_address, description=person.remark)
    for row in relationship_rows:
        add_edge(
            str(row.get("subject") or ""),
            str(row.get("related") or ""),
            str(row.get("relation") or "关系人"),
            str(row.get("wechat_id") or ""),
            ip_address=row.get("ip_address"),
            description=row.get("description"),
        )
    for index, left in enumerate(persons):
        left_clusters = [friend_cluster(row) for row in friends_of(left, friend_rows)]
        if not left_clusters:
            continue
        for right in persons[index + 1:]:
            right_clusters = [friend_cluster(row) for row in friends_of(right, friend_rows)]
            left_has_right = person_in_friend_list(right, left_clusters)
            right_has_left = person_in_friend_list(left, right_clusters)
            if not left_has_right and not right_has_left:
                continue
            mutual = left_has_right and right_has_left
            relation = relation_overrides.get((left.name, right.name)) or ("互为微信好友" if mutual else "微信好友")
            source_name, target_name = (left.name, right.name) if left_has_right else (right.name, left.name)
            add_edge(source_name, target_name, relation, kind="friend", mutual=mutual)

    sender_counts = Counter(item.sender_id for item in items if item.sender_id)
    for account, count in sender_counts.items():
        node_id = wechat_to_id.get(account)
        if node_id:
            nodes[node_id]["message_count"] = count

    if not subject_names and sender_counts:
        primary = sender_counts.most_common(1)[0][0]
        subject_names.add(primary)
        ensure_node(primary, primary, role="subject")
        for account, _ in sender_counts.most_common(21):
            if account != primary:
                add_edge(primary, account, "同群成员", account)

    degree = Counter()
    outgoing = Counter()
    for edge in edges:
        degree[edge["source"]] += 1
        degree[edge["target"]] += 1
        outgoing[edge["source"]] += 1
    for node_id, node in nodes.items():
        node["relation_count"] = degree[node_id]
        node["child_count"] = outgoing[node_id]

    groups: list[dict] = []
    grouped: dict[str, list[ChatItem]] = defaultdict(list)
    for item in items:
        if item.group_id:
            grouped[item.group_id].append(item)
    subject_ids = [node["id"] for node in nodes.values() if node["role"] == "subject"]
    for group_id, messages in grouped.items():
        activity = Counter(item.sender_id for item in messages if item.sender_id)
        active_accounts = [{"account": account, "message_count": count} for account, count in activity.most_common(5)]
        explicit_admin = next((
            match.group(1)
            for item in messages
            for match in [re.search(r"(?:群主|管理员)[：:\s]*(wxid_[A-Za-z0-9_-]+)", item.content)]
            if match
        ), None)
        inferred_admin = explicit_admin is None
        admin_account = explicit_admin or (active_accounts[0]["account"] if active_accounts else None)
        participating_subjects = [
            node_id for node_id in subject_ids
            if nodes[node_id].get("wechat_id") in activity
        ]
        groups.append({
            "group_id": group_id,
            "message_count": len(messages),
            "member_count": len(activity),
            "admin_account": admin_account,
            "admin_inferred": inferred_admin,
            "active_accounts": active_accounts,
            "participants": list(activity),
            "subject_ids": participating_subjects,
        })
    groups.sort(key=lambda group: group["message_count"], reverse=True)
    return {
        "subjects": subject_ids,
        "nodes": sorted(nodes.values(), key=lambda node: (node["role"] != "subject", -node["relation_count"], node["name"])),
        "edges": edges,
        "groups": groups,
        "relation_source": "人员关系数据" if relationship_rows or persons else "群聊共现推定",
    }


def format_person(person: MysqlPerson, title: str) -> str:
    parts = [f"{title}{person.name}"]
    fields = [
        ("性别", person.gender),
        ("身份证号", person.id_card),
        ("出生日期", person.birth_date),
        ("微信号", person.wechat_id),
        ("微信号别名", person.wechat_alias),
        ("原始微信号", person.original_wechat_id),
        ("微信昵称", person.wx_nickname),
        ("手机号", person.phone),
        ("快递地址", person.delivery_address),
        ("外卖地址", person.takeout_address),
        ("常用IP", person.ip_address),
        ("案件信息", person.case_info),
        ("关系备注", person.remark),
    ]
    details = [f"{label}：{value}" for label, value in fields if value]
    if details:
        parts.append("；".join(details))
    return "\n".join(parts)


def summarize_persons(persons: list[MysqlPerson]) -> tuple[list[str], list[MysqlPerson], list[MysqlPerson]]:
    if not persons:
        return ["本案还没有人员信息。请在案件中导入带姓名的人员表。"], [], []
    subjects = subjects_of(persons)
    related = related_of(persons, subjects)
    lines = [format_person(person, "【本人】") for person in subjects]
    if related:
        lines.append(f"【关系人】共 {len(related)} 人")
        lines.extend(format_person(person, "【关系人】") for person in related)
    else:
        lines.append("【关系人】未登记其他关系人")
    return lines, subjects, related


def lexical_hit(question: str, content: str) -> bool:
    query = question.strip().lower()
    if not query:
        return False
    if query in content.lower():
        return True
    tokens = [token for token in re.findall(r"[\u4e00-\u9fff]{2,}|[a-z0-9_]{2,}", query) if token]
    return any(token in content.lower() for token in tokens)


def search_keyword_items(question: str, items: list[ChatItem], min_score: float = 0.42) -> list[tuple[ChatItem, float]]:
    """BGE-M3 关键词/语义检索，命中即输出，不做 Top10 截断。"""
    query_vector: list[float] | None = None
    if settings.embedding_backend == "bge":
        try:
            query_vector = encode_bge([question])[0]
        except Exception:
            query_vector = None

    scored: list[tuple[ChatItem, float]] = []
    for item in items:
        exact = lexical_hit(question, item.content)
        similarity = cosine(query_vector, item.embedding) if query_vector and item.embedding else 0.0
        if exact or similarity >= min_score:
            score = max(similarity, 1.0 if exact and question.strip().lower() in item.content.lower() else similarity)
            if exact:
                score = max(score, 0.99 if question.strip().lower() in item.content.lower() else 0.7)
            scored.append((item, score))
    scored.sort(key=lambda pair: pair[1], reverse=True)
    return scored


def evidence_harness(
    question: str,
    items: list[ChatItem],
    limit: int = 10,
    before: int = 2,
    after: int = 2,
) -> list[SearchHit]:
    """在任意已规范化数据源上先召回明确原文，再扩展同群上下文。"""
    normalized = question.strip().lower()
    existence = re.search(r"(?:有没有|是否|有无)(?:有)?([\u4e00-\u9fff]{1,4})", normalized)
    reduced = normalized
    for word in QUESTION_WORDS:
        reduced = reduced.replace(word, "")
    terms = set(re.findall(r"wxid_[a-z0-9]+|(?:\d{1,3}\.){3}\d{1,3}|[a-z0-9_]{2,}|[\u4e00-\u9fff]{2,}", reduced))
    focus_term = existence.group(1).rstrip("吗么呢") if existence else ""
    if focus_term:
        terms.add(focus_term)
    for phrase in tuple(terms):
        if re.fullmatch(r"[\u4e00-\u9fff]{2,}", phrase):
            terms.update(phrase[index:index + 2] for index in range(len(phrase) - 1))
            if len(phrase) <= 3:
                terms.update(char for char in phrase if char not in "有的是了在")

    semantic = {item.message_id: score for item, score in search_keyword_items(question, items)}
    scored: list[tuple[ChatItem, float]] = []
    for item in items:
        fields = " ".join(value.lower() for value in (item.sender_id, item.group_id, item.ip_address) if value)
        content = item.content.lower()
        field_hits = sum(term in fields for term in terms)
        content_hits = sum(term in content for term in terms)
        score = semantic.get(item.message_id, 0.0) + field_hits * 4.0 + content_hits * 2.0
        if focus_term and focus_term in content:
            score += 8.0
        if normalized in content:
            score += 6.0
        if score > 0:
            scored.append((item, score))
    ranked = sorted(scored, key=lambda pair: pair[1], reverse=True)
    cutoff = ranked[0][1] * 0.5 if ranked else 0.0
    anchors: list[tuple[ChatItem, float]] = []
    seen_content: set[tuple[str | None, str]] = set()
    for item, score in ranked:
        content_key = (item.group_id, item.content.strip().lower())
        if score < cutoff or content_key in seen_content:
            continue
        anchors.append((item, score))
        seen_content.add(content_key)
        if len(anchors) == limit:
            break
    selected = {item.message_id: (item, score) for item, score in anchors}

    grouped: dict[tuple[object, str | None], list[ChatItem]] = defaultdict(list)
    for item in items:
        grouped[(item.source_file_id, item.group_id)].append(item)
    for rows in grouped.values():
        rows.sort(key=lambda item: (item.sent_at is None, item.sent_at, item.source_row))
    for anchor, score in anchors:
        rows = grouped[(anchor.source_file_id, anchor.group_id)]
        position = next(index for index, item in enumerate(rows) if item.message_id == anchor.message_id)
        for distance, item in enumerate(rows[max(0, position - before):position + after + 1]):
            context_score = max(0.01, score - abs(max(0, position - before) + distance - position) * 0.01)
            if item.message_id not in selected:
                selected[item.message_id] = (item, context_score)
    return [to_search_hit(item, score) for item, score in sorted(selected.values(), key=lambda pair: pair[1], reverse=True)]


def chats_to_items(chats: list[MysqlChat]) -> list[ChatItem]:
    items: list[ChatItem] = []
    for chat in chats:
        source_key = f"mysql:{chat.source_file or 'import'}"
        items.append(
            ChatItem(
                message_id=chat.message_id,
                source_file_id=uuid.uuid5(uuid.NAMESPACE_URL, source_key),
                source_row=chat.source_row or 0,
                sent_at=chat.sent_at,
                group_id=chat.group_id,
                sender_id=chat.wechat_id,
                content=chat.content,
                ip_address=chat.ip_address,
                embedding=chat.embedding,
            )
        )
    return items


def to_search_hit(item: ChatItem, score: float) -> SearchHit:
    return SearchHit(
        message_id=item.message_id,
        source_file_id=item.source_file_id,
        source_row=item.source_row,
        sent_at=item.sent_at,
        group_id=item.group_id,
        sender_id=item.sender_id,
        content=item.content,
        ip_address=item.ip_address,
        score=round(float(score), 6),
    )


def build_group_profiles(
    items: list[ChatItem],
    subjects: list[MysqlPerson],
    related: list[MysqlPerson],
) -> list[GroupProfile]:
    subject_wxids = {person.wechat_id for person in subjects if person.wechat_id}
    related_map = {person.wechat_id: person for person in related if person.wechat_id}
    if not subject_wxids:
        return []

    by_group: dict[str, list[ChatItem]] = defaultdict(list)
    for item in items:
        if not item.group_id:
            continue
        by_group[item.group_id].append(item)

    subject_groups = {
        group_id
        for group_id, rows in by_group.items()
        if any(row.sender_id in subject_wxids for row in rows)
    }
    profiles: list[GroupProfile] = []
    for group_id in subject_groups:
        rows = by_group[group_id]
        senders = [row.sender_id for row in rows if row.sender_id]
        related_present = [related_map[wxid] for wxid in dict.fromkeys(senders) if wxid in related_map]
        ordered = sorted(rows, key=lambda item: (item.sent_at is None, item.sent_at, item.source_row))
        messages = [
            {
                "sender_id": row.sender_id,
                "sent_at": row.sent_at.isoformat(sep=" ", timespec="seconds") if row.sent_at else None,
                "content": (row.content or "").strip(),
            }
            for row in ordered
            if (row.content or "").strip()
        ]
        profiles.append(
            GroupProfile(
                group_id=group_id,
                member_count=len(set(senders)),
                members=sorted(set(senders)),
                message_count=len(rows),
                subject_speak_count=sum(1 for row in rows if row.sender_id in subject_wxids),
                related_count=len(related_present),
                related_names=[f"{person.name}（{person.remark or person.wechat_id}）" for person in related_present],
                messages=messages,
            )
        )
    return sort_groups_by_activity(profiles)


def sort_groups_by_activity(profiles: list[GroupProfile]) -> list[GroupProfile]:
    profiles.sort(key=lambda item: (item.subject_speak_count, item.related_count, item.message_count), reverse=True)
    return profiles


def classify_groups_heuristic(profiles: list[GroupProfile]) -> list[GroupProfile]:
    ranked = sort_groups_by_activity(profiles)
    for index, profile in enumerate(ranked, start=1):
        profile.reason = (
            f"排序第 {index}：本人发言 {profile.subject_speak_count} 次，"
            f"关系人 {profile.related_count} 人"
        )
        profile.category = "待研判"
    return ranked


def group_payload(profiles: list[GroupProfile]) -> list[dict]:
    ranked = sort_groups_by_activity(list(profiles))
    return [
        {
            "rank": index,
            "group_id": profile.group_id,
            "member_count": profile.member_count,
            "message_count": profile.message_count,
            "subject_speak_count": profile.subject_speak_count,
            "related_count": profile.related_count,
            "related_names": profile.related_names,
            "chats": profile.messages,
        }
        for index, profile in enumerate(ranked, start=1)
    ]


def apply_group_judgement(profiles: list[GroupProfile], judgements: list[dict]) -> list[GroupProfile]:
    mapping = {item.get("group_id"): item for item in judgements if item.get("group_id")}
    for profile in profiles:
        item = mapping.get(profile.group_id)
        if not item:
            continue
        category = item.get("category")
        if category in ("关键群", "普通家人群"):
            profile.category = category
        if item.get("reason"):
            profile.reason = str(item["reason"])
    return profiles


def format_group_overview(profiles: list[GroupProfile], subjects: list[MysqlPerson]) -> str:
    wxids = "、".join(sorted({person.wechat_id for person in subjects if person.wechat_id})) or "未登记"
    if not profiles:
        return f"【群聊】本人微信号（{wxids}）在已导入聊天中未发现所在群。"
    lines = [f"【群聊】本人微信号 {wxids} 共出现在 {len(profiles)} 个群"]
    for profile in profiles:
        related = "、".join(profile.related_names) if profile.related_names else "无"
        lines.append(
            f"群 {profile.group_id}：成员 {profile.member_count} 人，"
            f"发言 {profile.message_count} 条（本人 {profile.subject_speak_count} 条），"
            f"关系人 {profile.related_count} 人（{related}）"
        )
    return "\n".join(lines)


def load_case_materials(
    case_number: str,
    sqlite_items: list[ChatItem],
    case_name: str | None = None,
    persons: list[MysqlPerson] | None = None,
) -> tuple[list[MysqlPerson], list[ChatItem]]:
    return list(persons or []), list(sqlite_items)


def person_brief(person: MysqlPerson) -> dict:
    return {
        "name": person.name,
        "gender": person.gender,
        "id_card": person.id_card,
        "birth_date": person.birth_date,
        "wechat_id": person.wechat_id,
        "wechat_alias": person.wechat_alias,
        "original_wechat_id": person.original_wechat_id,
        "wx_nickname": person.wx_nickname,
        "phone": person.phone,
        "delivery_address": person.delivery_address,
        "takeout_address": person.takeout_address,
        "ip_address": person.ip_address,
        "case_info": person.case_info,
        "remark": person.remark,
    }


def wants_full_profile(question: str) -> bool:
    text = (question or "").strip()
    if not text:
        return True
    return bool(PROFILE_QUESTION.search(text))


def build_llm_briefing(persons: list[MysqlPerson], profiles: list[GroupProfile], question: str = "") -> dict:
    """整理本人、关系人、群统计，只作为模型输入，不直接展示。"""
    subjects = subjects_of(persons) if persons else []
    related = related_of(persons, subjects) if persons else []
    if wants_full_profile(question):
        task = [
            "先总结本人全部个人信息，再分别总结各关系人信息",
            "群已按本人发言次数、关系人数量排序",
            "必须阅读每个群 chats 里的群聊原文，再结合关系人构成，判断该群疑似是什么群（如家庭群、同事群、朋友群、工作群、其他），写成 inference，并引用具体发言作为依据",
            "群类型 inference 的 statement 写成一句：群{群名称或群ID}，疑似为{类型}，成员包括{成员}，聊天内容涉及{聊天要点}",
            "不要自行检索违法犯罪关键词，也不要编造材料中没有的事实",
        ]
    else:
        task = [
            "只回答当前 question，不要把本人、关系人或全部群类型整份复述一遍",
            "仅在问题确实问到某人、某关系或某群时，才摘用 briefing 中对应条目",
            "不要自行检索违法犯罪关键词，也不要编造材料中没有的事实",
        ]
    return {
        "subjects": [person_brief(person) for person in subjects],
        "related_persons": [person_brief(person) for person in related],
        "groups": group_payload(profiles),
        "task": task,
    }
