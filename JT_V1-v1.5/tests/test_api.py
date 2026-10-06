import json
from io import BytesIO
from pathlib import Path
from uuid import UUID

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.db import SessionLocal
from app.main import app, parse_case_focus_groups, parse_inferred_group_types
from app.models import Analysis
from app.config import settings
from app.services import stable_message_id


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "聊天记录"
    sheet.append(["序号", "时间", "群ID", "微信号", "聊天内容", "IP登录地址"])
    sheet.append([1, "2026-09-01 10:00:00", "g1", "wxid_l3q7j5v9", "蓝盒还在原处吗", "127.0.0.1"])
    sheet.append([2, "2026-09-01 10:02:00", "g1", "u2", "昨晚没动，等通知", "127.0.0.2"])
    relations = workbook.create_sheet("人员关系")
    relations.append(["主体", "关系人", "关系类型", "微信号", "IP地址", "关系说明"])
    relations.append(["周凯", "刘桂香", "母亲", "wxid_liuguixiang", "127.0.0.3", "周凯的母亲"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def conversation_only_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "聊天记录"
    sheet.append(["序号", "时间", "群ID", "微信号", "聊天内容", "IP登录地址"])
    sheet.append([1, "2026-09-02 11:00:00", "g-private", "wxid_private", "对话专属句子", "10.0.0.9"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def context_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "聊天记录"
    sheet.append(["序号", "时间", "群ID", "微信号", "聊天内容", "IP登录地址"])
    sheet.append([1, "2026-09-01 09:58:00", "g1", "u1", "前文消息", "127.0.0.1"])
    sheet.append([2, "2026-09-01 10:00:00", "g1", "u2", "目标消息", "127.0.0.2"])
    sheet.append([3, "2026-09-01 10:02:00", "g1", "u1", "后文消息", "127.0.0.1"])
    sheet.append([4, "2026-09-01 10:01:00", "g2", "u3", "其他群消息", "127.0.0.3"])
    sheet.append([5, "2026-09-01 10:30:00", "g1", "u2", "超出时间邻域", "127.0.0.2"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_focus_keywords_mark_groups_after_material_analysis():
    with TestClient(app) as client:
        created = client.post("/cases", json={
            "name": "关键词案",
            "code": "FOCUS-001",
            "category": "诈骗",
            "handler": "林警官",
            "unit": "某分局",
            "summary": "嫌疑人通过群聊约定取货",
            "keywords": "目标消息、无关词",
        })
        assert created.status_code == 201
        assert created.json()["keywords"] == "目标消息、无关词"
        case_id = created.json()["id"]
        imported = client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("chat.xlsx", context_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert imported.status_code == 201
        before = client.get(f"/cases/{case_id}/groups")
        assert before.status_code == 200
        assert all(item["focus"] is False for item in before.json())
        conversation = client.post(f"/cases/{case_id}/conversations", json={"title": "研判"}).json()
        scoped = client.get(f"/cases/{case_id}/groups", params={"conversation_id": conversation["id"]})
        assert all(item["focus"] is False for item in scoped.json())
        client.patch(f"/conversations/{conversation['id']}", json={"analysis_enabled": True})
        marked = client.get(f"/cases/{case_id}/groups", params={"conversation_id": conversation["id"]}).json()
        by_id = {item["group_id"]: item for item in marked}
        assert by_id["g1"]["focus"] is True
        assert by_id["g1"]["focus_terms"] == ["目标消息"]
        assert by_id["g2"]["focus"] is False
        assert marked[0]["group_id"] == "g1"


def test_parse_case_focus_groups_keeps_possible_matches_only():
    found = parse_case_focus_groups([
        "群g-both，重点群，已导入人员包括佳宁、高磊",
        "群g-maybe，重点群，可能相关，依据：约定取货",
    ])
    assert found == {"g-maybe": ["可能相关"]}


def test_focus_group_matches_either_condition_or_both():
    sheet = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    with TestClient(app) as client:
        case_id = client.post("/cases", json={
            "name": "或条件案",
            "code": "FOCUS-OR-001",
            "category": "敲诈勒索",
            "summary": "团伙通过群聊约定取货",
            "keywords": "看货",
        }).json()["id"]
        files = [
            ("高磊_gaolei98766877.xlsx", gang_friend_bytes([("wxid_7k2m3n8p", "wxid_7k2m3n8p", "佳宁", "佳宁")])),
            ("佳宁_wxid_7k2m3n8p.xlsx", gang_friend_bytes([("gaolei98766877", "gaolei98766877", "高磊", "高磊")])),
            ("记录_群聊.xlsx", gang_chat_bytes([
                ("g-both", "gaolei98766877", "高磊", "今晚看货"),
                ("g-both", "wxid_7k2m3n8p", "佳宁", "地点我来定"),
                ("g-people", "gaolei98766877", "高磊", "随便聊聊"),
                ("g-people", "wxid_7k2m3n8p", "佳宁", "知道了"),
                ("g-word", "gaolei98766877", "高磊", "单独看货"),
                ("g-none", "gaolei98766877", "高磊", "今天天气不错"),
                ("g-maybe", "wxid_other", "路人", "约定取货"),
            ])),
        ]
        for name, payload in files:
            imported = client.post(f"/cases/{case_id}/sources", files={"file": (name, payload, sheet)})
            assert imported.status_code == 201, imported.text
        conversation = client.post(f"/cases/{case_id}/conversations", json={"title": "研判"}).json()
        client.patch(f"/conversations/{conversation['id']}", json={"analysis_enabled": True})
        with SessionLocal() as session:
            session.add(Analysis(
                case_id=UUID(case_id),
                conversation_id=UUID(conversation["id"]),
                question="请根据案件信息判断群类型",
                result={"claims": [{"type": "inference", "statement": "群g-maybe，重点群，可能相关，依据：约定取货"}]},
                engine="test",
            ))
            session.commit()
        marked = {
            item["group_id"]: item
            for item in client.get(f"/cases/{case_id}/groups", params={"conversation_id": conversation["id"]}).json()
        }
        assert marked["g-both"]["focus"] is True
        assert set(marked["g-both"]["focus_people"]) == {"高磊", "佳宁"}
        assert marked["g-both"]["focus_terms"] == ["看货"]
        assert marked["g-people"]["focus"] is True
        assert marked["g-people"]["focus_terms"] == []
        assert marked["g-word"]["focus"] is True
        assert marked["g-word"]["focus_people"] == []
        assert marked["g-word"]["focus_terms"] == ["看货"]
        assert marked["g-none"]["focus"] is False
        assert marked["g-maybe"]["focus"] is True
        assert marked["g-maybe"]["focus_case"] == ["可能相关"]
        assert marked["g-maybe"]["focus_people"] == []


def gang_friend_bytes(rows: list[tuple[str, str, str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "团伙成员好友列表"
    sheet.append(["好友微信号", "好友原始账号", "呢称", "姓名", "好友备注"])
    for wechat_id, original, nickname, real_name in rows:
        sheet.append([wechat_id, original, nickname, real_name, ""])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def gang_chat_bytes(rows: list[tuple[str, str, str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "群聊"
    sheet.append(["群号(GROUP)", "发送方账号ID(SENDID)", "发送方呢称", "消息发出时间(CAPTURE_TIME)", "聊天内容(CONTENT)", "发送方IP(SENDIP)"])
    for group_id, sender_id, nickname, content in rows:
        sheet.append([group_id, sender_id, nickname, "2026-06-20 08:08:13", content, "127.0.0.8"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_uploaded_friend_and_chat_formats_mark_partial_overlap():
    sheet = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "格式案", "code": "FORMAT-001"}).json()["id"]
        files = [
            ("高磊_gaolei98766877.xlsx", gang_friend_bytes([
                ("", "wxid_ziwei", "紫薇", "紫薇"),
                ("wxid_7k2m3n8p", "wxid_7k2m3n8p", "佳宁", "佳宁"),
            ])),
            ("佳宁_wxid_7k2m3n8p.xlsx", gang_friend_bytes([
                ("", "wxid_ziwei", "紫薇", "紫薇"),
                ("gaolei98766877", "gaolei98766877", "高磊", "高磊"),
            ])),
            ("石鎏_wxid_3w5x7y9z.xlsx", gang_friend_bytes([
                ("gaolei98766877", "gaolei98766877", "高磊", "高磊"),
            ])),
            ("高磊_群聊.xlsx", gang_chat_bytes([
                ("g-two", "gaolei98766877", "高磊", "一起看货"),
                ("g-two", "wxid_7k2m3n8p", "佳宁", "明天到"),
                ("g-one", "gaolei98766877", "高磊", "只有我"),
                ("g-out", "wxid_ziwei", "紫薇", "外面的人"),
                ("g-out", "wxid_other", "路人", "无关"),
            ])),
            ("佳宁_群聊.xlsx", gang_chat_bytes([
                ("g-two", "wxid_7k2m3n8p", "佳宁", "群里见"),
            ])),
        ]
        for name, payload in files:
            imported = client.post(f"/cases/{case_id}/sources", files={"file": (name, payload, sheet)})
            assert imported.status_code == 201, imported.text
        graph = client.get(f"/cases/{case_id}/graph").json()
        owners = {item["name"]: item for item in graph["nodes"]}
        assert owners["高磊"]["description"] == "本人"
        assert owners["高磊"]["wechat_id"] == "gaolei98766877"
        assert owners["佳宁"]["wechat_id"] == "wxid_7k2m3n8p"
        assert "紫薇" not in owners
        conversation = client.post(f"/cases/{case_id}/conversations", json={"title": "研判"}).json()
        before = client.get(f"/cases/{case_id}/groups", params={"conversation_id": conversation["id"]}).json()
        assert all(item["focus"] is False for item in before)
        owned = {item["group_id"]: item["wechat_owners"] for item in before}
        assert owned["g-two"] == ["佳宁的微信 wxid_7k2m3n8p", "高磊的微信 gaolei98766877"]
        assert owned["g-one"] == ["高磊的微信 gaolei98766877"]
        assert owned["g-out"] == ["高磊的微信 gaolei98766877"]
        client.patch(f"/conversations/{conversation['id']}", json={"analysis_enabled": True})
        marked = {item["group_id"]: item for item in client.get(f"/cases/{case_id}/groups", params={"conversation_id": conversation["id"]}).json()}
        assert marked["g-two"]["focus"] is True
        assert set(marked["g-two"]["focus_people"]) == {"高磊", "佳宁"}
        assert marked["g-one"]["focus"] is False
        assert marked["g-out"]["focus"] is False
        screened = client.post(f"/cases/{case_id}/friend-screen", json={"names": ["高磊", "佳宁", "石鎏"]}).json()
        hit = next(item for item in screened["common_friends"] if item["wechat_id"] == "wxid_ziwei")
        assert hit["shared_count"] == 2
        assert set(hit["shared_by"]) == {"高磊", "佳宁"}
        analysed = client.post(f"/cases/{case_id}/analysis", json={
            "question": "请总结人员信息并判断群类型",
            "conversation_id": conversation["id"],
            "model_mode": "local",
            "limit": 5,
        })
        assert analysed.status_code == 201
        statements = [claim["statement"] for claim in analysed.json()["claims"]]
        assert any(item.startswith("群g-two，重点群，已导入人员包括") and "佳宁" in item and "高磊" in item for item in statements)
        assert not any(item.startswith("群g-one") or item.startswith("群g-out") for item in statements)


def test_parse_inferred_group_types():
    found = parse_inferred_group_types([
        "群78045319，疑似为朋友群，成员包括周凯、郑浩、吴磊",
        "群69035712，疑似为同事群/工作群，成员包括周凯、赵立国",
        "群g1 疑似是家庭群",
        "周凯，男，身份证号330881198805125172",
    ])
    assert found["78045319"] == "朋友群"
    assert found["69035712"] == "同事群/工作群"
    assert found["g1"] == "家庭群"


def test_stable_message_id():
    assert stable_message_id("abc", "聊天记录", 2) == stable_message_id("abc", "聊天记录", 2)
    assert stable_message_id("abc", "聊天记录", 2) != stable_message_id("abc", "聊天记录", 3)
    assert stable_message_id("abc", "聊天记录", 2, "CASE-A") != stable_message_id("abc", "聊天记录", 2, "CASE-B")


def test_frontend_and_case_listing():
    with TestClient(app) as client:
        frontend = client.get("/ui").text
        assert "酒桶研判系统" in frontend
        assert "群聊原始数据" in frontend
        assert "案件项目" in frontend
        assert 'data-case-search' in frontend
        assert 'data-import-mode' in frontend
        assert '导入案件材料' in frontend
        assert '材料智能分析' in frontend
        assert '对话材料' in frontend
        assert '案件材料' in frontend
        assert '简要案情' in frontend
        assert '重点关键词' in frontend
        assert 'data-case-keywords-input' in frontend
        assert '承办人' in frontend
        assert 'data-api-preset-save' in frontend
        assert '人工识别' in frontend
        assert 'data-file-input multiple' in frontend
        assert '可一次选择多个表格' in frontend
        assert 'data-import-mode="ai" hidden' in frontend
        assert 'data-language-toggle hidden' in frontend
        assert 'data-review-dialog' in frontend
        assert '保存在当前案件' in frontend
        assert 'data-resizer="sidebar"' in frontend
        assert 'data-graph-add-subject' in frontend
        assert 'data-graph-delete-subjects' in frontend
        assert '微信好友表' in frontend
        assert client.get("/theme.config.css").status_code == 200
        script = client.get("/app.js").text
        assert "group-owner-mark" in script
        assert "来源未识别" in script
        assert "relation_count ?? 0" in script
        assert "notes: value" in script
        assert "item.wechat_id || item.wechat_alias || item.original_wechat_id" in script
        assert "group-owner-mark" in client.get("/theme.config.css").text

        legacy_frontend = client.get("/ui/legacy").text
        assert "群聊分析研判系统" in legacy_frontend
        assert 'id="caseManageBtn"' in legacy_frontend
        assert 'id="keywordSearchBtn"' in legacy_frontend
        assert '/keyword-search' in legacy_frontend
        assert 'id="caseTable"' in legacy_frontend
        assert "无法连接后端" in legacy_frontend
        assert "location.protocol==='file:'" in legacy_frontend
        assert client.get("/cases").json() == []


def test_case_notes_round_trip():
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "笔记案", "code": "NOTE-001"}).json()["id"]
        updated = client.patch(f"/cases/{case_id}", json={"notes": "换浏览器也能看到"}).json()
        assert updated["notes"] == "换浏览器也能看到"
        listed = client.get("/cases").json()
        assert next(item["notes"] for item in listed if item["id"] == case_id) == "换浏览器也能看到"


def test_rejects_invalid_workbook():
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "测试案件", "code": "INVALID-001"}).json()["id"]
        response = client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("fake.xlsx", b"not a workbook", "application/octet-stream")},
        )
        assert response.status_code == 422
        assert response.json()["detail"] == "文件不是有效的 xlsx 工作簿"


def test_end_to_end():
    with TestClient(app) as client:
        assert client.get("/health").json() == {"status": "ok"}

        created = client.post("/cases", json={
            "name": "测试案件",
            "code": "CASE-001",
            "category": "诈骗",
            "handler": "林警官",
            "unit": "某分局",
            "summary": "简要案情测试",
        })
        assert created.status_code == 201
        assert created.json()["handler"] == "林警官"
        case_id = created.json()["id"]
        payload = workbook_bytes()

        imported = client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("chat.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert imported.status_code == 201
        assert imported.json()["imported_messages"] == 2
        assert imported.json()["imported_persons"] == 2

        stats = client.get(f"/cases/{case_id}/stats")
        assert stats.json() == {
            "source_files": 1,
            "case_source_files": 1,
            "messages": 2,
            "analyses": 0,
            "reviews": 0,
        }
        assert imported.json()["scope"] == "case"

        duplicate = client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("chat.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert duplicate.json()["duplicate"] is True

        results = client.get(f"/cases/{case_id}/search", params={"q": "蓝盒还在原处吗"})
        assert results.status_code == 200
        assert results.json()[0]["content"] == "蓝盒还在原处吗"
        assert results.json()[0]["source_row"] == 2

        groups = client.get(f"/cases/{case_id}/groups")
        assert groups.status_code == 200
        assert groups.json()[0]["group_id"] == "g1"
        assert groups.json()[0]["message_count"] == 2
        graph = client.get(f"/cases/{case_id}/graph")
        assert graph.status_code == 200
        assert graph.json()["relation_source"] == "人员关系数据"
        assert [node["name"] for node in graph.json()["nodes"][:2]] == ["周凯", "刘桂香"]
        assert graph.json()["edges"] == [{
            "source": "name:周凯",
            "target": "wx:wxid_liuguixiang",
            "relation": "母亲",
        }]
        assert graph.json()["groups"][0]["admin_account"] == "wxid_l3q7j5v9"
        assert graph.json()["groups"][0]["admin_inferred"] is True
        group_messages = client.get(
            f"/cases/{case_id}/group-messages", params={"group_id": "g1"}
        )
        assert [item["content"] for item in group_messages.json()] == [
            "蓝盒还在原处吗",
            "昨晚没动，等通知",
        ]

        conversation = client.post(
            f"/cases/{case_id}/conversations", json={"title": "蓝盒研判"}
        )
        assert conversation.status_code == 201
        conversation_id = conversation.json()["id"]
        assert conversation.json()["analysis_enabled"] is False
        assert client.get(f"/cases/{case_id}/conversations").json()[0]["title"] == "蓝盒研判"

        overview = client.post(
            f"/cases/{case_id}/analysis",
            json={"question": "这个案件上传了哪些文件", "conversation_id": conversation_id, "limit": 10},
        )
        assert overview.status_code == 201
        assert overview.json()["engine"] == "case-overview"
        overview_text = " ".join(item["statement"] for item in overview.json()["claims"])
        assert "林警官" in overview_text
        assert "chat.xlsx" in overview_text
        assert overview.json()["evidence"] == []

        enabled = client.patch(
            f"/conversations/{conversation_id}", json={"analysis_enabled": True}
        )
        assert enabled.status_code == 200
        assert enabled.json()["analysis_enabled"] is True

        analysis = client.post(
            f"/cases/{case_id}/analysis",
            json={"question": "蓝盒还在原处吗", "conversation_id": conversation_id, "limit": 10},
        )
        assert analysis.status_code == 201
        assert analysis.json()["review_status"] == "待复核"
        assert analysis.json()["engine"] == "awaiting-llm"
        assert analysis.json()["evidence"][0]["content"] == "蓝盒还在原处吗"
        turns = client.get(f"/conversations/{conversation_id}/turns")
        assert turns.status_code == 200
        engines = {item["question"]: item["result"]["engine"] for item in turns.json()}
        assert engines["这个案件上传了哪些文件"] == "case-overview"
        assert engines["蓝盒还在原处吗"] == "awaiting-llm"

        summary = client.post(
            f"/cases/{case_id}/analysis",
            json={"question": "请总结人员信息并判断群类型", "conversation_id": conversation_id, "limit": 10},
        )
        assert summary.status_code == 201
        assert summary.json()["engine"] == "awaiting-llm"
        assert summary.json()["judgement_level"] != "证据不足"
        assert summary.json()["evidence"]
        assert summary.json()["subjects"]

        keyword_hits = client.get(f"/cases/{case_id}/keyword-search", params={"q": "蓝盒"})
        assert keyword_hits.status_code == 200
        assert keyword_hits.json()[0]["content"] == "蓝盒还在原处吗"

        updated = client.patch(
            f"/conversations/{conversation_id}",
            json={"title": "置顶蓝盒研判", "pinned": True, "unread": True, "section": "待复核"},
        )
        assert updated.status_code == 200
        assert updated.json()["pinned"] is True
        assert updated.json()["section"] == "待复核"

        forked = client.post(f"/conversations/{conversation_id}/fork", json={})
        assert forked.status_code == 201
        fork_id = forked.json()["id"]
        assert forked.json()["analysis_enabled"] is True
        assert len(client.get(f"/conversations/{fork_id}/turns").json()) == 3

        other_case_id = client.post(
            "/cases", json={"name": "其他案件", "code": "CASE-002"}
        ).json()["id"]
        blocked_move = client.patch(
            f"/conversations/{conversation_id}", json={"case_id": other_case_id}
        )
        assert blocked_move.status_code == 409

        ip_analysis = client.post(
            f"/cases/{case_id}/analysis",
            json={
                "question": "wxid_l3q7j5v9的登录地址是什么",
                "conversation_id": conversation_id,
                "limit": 6,
            },
        )
        assert ip_analysis.status_code == 201
        assert ip_analysis.json()["engine"] == "structured-field-lookup"
        assert ip_analysis.json()["subjects"] == ["wxid_l3q7j5v9"]
        assert "127.0.0.1" in ip_analysis.json()["claims"][0]["statement"]
        assert ip_analysis.json()["evidence"][0]["ip_address"] == "127.0.0.1"

        followup = client.post(
            f"/cases/{case_id}/analysis",
            json={
                "question": "他的登录地址呢",
                "conversation_id": conversation_id,
                "limit": 6,
            },
        )
        assert followup.status_code == 201
        assert followup.json()["engine"] == "structured-field-lookup"
        assert "127.0.0.1" in followup.json()["claims"][0]["statement"]

        queued = client.post(
            f"/cases/{case_id}/analysis-tasks",
            json={
                "question": "wxid_l3q7j5v9的登录地址是什么",
                "conversation_id": conversation_id,
                "limit": 6,
            },
        )
        assert queued.status_code == 202
        completed = client.get(f"/analysis-tasks/{queued.json()['id']}")
        assert completed.json()["status"] == "succeeded"
        assert completed.json()["result"]["engine"] == "structured-field-lookup"

        api_unconfigured = client.post(
            f"/cases/{case_id}/analysis",
            json={"question": "蓝盒是否移动", "model_mode": "api", "limit": 2},
        )
        assert api_unconfigured.status_code == 503
        assert "API 地址" in api_unconfigured.json()["detail"]
        assert "模型名" in api_unconfigured.json()["detail"]

        analysis_id = analysis.json()["id"]
        invalid_revision = client.post(
            f"/analyses/{analysis_id}/reviews",
            json={"decision": "修改", "reviewer": "林警官", "comment": "需要修改"},
        )
        assert invalid_revision.status_code == 422

        reviewed = client.post(
            f"/analyses/{analysis_id}/reviews",
            json={"decision": "采纳", "reviewer": "林警官", "comment": "证据引用无误"},
        )
        assert reviewed.status_code == 201
        assert reviewed.json()["decision"] == "采纳"

        evidence_id = client.get(f"/cases/{case_id}/keyword-search", params={"q": "蓝盒"}).json()[0]["message_id"]
        modified = client.post(
            f"/analyses/{analysis_id}/reviews",
            json={
                "decision": "修改",
                "reviewer": "周警官",
                "comment": "降低结论等级",
                "revised_result": {
                    "judgement_level": "证据不足",
                    "claims": [{
                        "type": "fact",
                        "statement": "存在相关原始消息",
                        "evidence_refs": [evidence_id],
                    }],
                    "counter_evidence_refs": [],
                    "alternatives": [],
                    "confidence": "低",
                },
            },
        )
        assert modified.status_code == 201

        history = client.get(f"/analyses/{analysis_id}/reviews")
        assert history.status_code == 200
        assert [item["decision"] for item in history.json()] == ["采纳", "修改"]

        removed = client.delete(f"/conversations/{conversation_id}")
        assert removed.status_code == 204
        assert client.delete(f"/conversations/{fork_id}").status_code == 204
        assert client.get(f"/cases/{case_id}/conversations").json() == []
        assert client.get(f"/analyses/{analysis_id}/reviews").status_code == 404


def test_context_and_time_neighbors_stay_in_group():
    with TestClient(app) as client:
        case_id = client.post(
            "/cases", json={"name": "上下文测试", "code": "CTX-001"}
        ).json()["id"]
        imported = client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("context.xlsx", context_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert imported.json()["imported_messages"] == 5

        anchor = client.get(
            f"/cases/{case_id}/search", params={"q": "目标消息", "limit": 1}
        ).json()[0]

        context = client.get(
            f"/cases/{case_id}/messages/{anchor['message_id']}/context",
            params={"before": 1, "after": 1},
        )
        assert context.status_code == 200
        assert [item["content"] for item in context.json()] == ["前文消息", "目标消息", "后文消息"]
        assert [item["relation"] for item in context.json()] == ["before", "anchor", "after"]

        neighbors = client.get(
            f"/cases/{case_id}/messages/{anchor['message_id']}/time-neighbors",
            params={"minutes": 5},
        )
        assert neighbors.status_code == 200
        assert [item["content"] for item in neighbors.json()] == ["前文消息", "目标消息", "后文消息"]
        assert [item["seconds_from_anchor"] for item in neighbors.json()] == [-120, 0, 120]

        options = client.get(f"/cases/{case_id}/filter-options").json()
        assert options["start_date"] == "2026-09-01"
        assert options["groups"] == ["g1", "g2"]
        assert options["senders"] == ["u1", "u2", "u3"]

        filtered = client.get(
            f"/cases/{case_id}/search",
            params={"q": "消息", "group_id": "g2", "sender_id": "u3"},
        )
        assert [item["content"] for item in filtered.json()] == ["其他群消息"]

        conflict = client.get(
            f"/cases/{case_id}/search",
            params={"q": "消息", "start_date": "2026-09-02", "end_date": "2026-09-01"},
        )
        assert conflict.status_code == 422
        assert conflict.json()["detail"] == "开始日期不能晚于结束日期"

        no_evidence = client.post(
            f"/cases/{case_id}/analysis",
            json={"question": "目标消息在哪里", "sender_id": "does-not-exist"},
        )
        assert no_evidence.status_code == 201
        assert no_evidence.json()["engine"] == "no-evidence"
        assert no_evidence.json()["judgement_level"] == "证据不足"


def test_case_data_management_and_audit():
    with TestClient(app) as client:
        case_id = client.post(
            "/cases", json={"name": "数据管理测试", "code": "DATA-001"}
        ).json()["id"]
        client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("chat.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        sources = client.get(f"/cases/{case_id}/sources")
        assert sources.json()[0]["filename"] == "chat.xlsx"
        audit = client.get(f"/cases/{case_id}/audit-logs")
        assert {item["action"] for item in audit.json()} >= {"case.create", "source.import"}
        exported = client.get(f"/cases/{case_id}/export")
        assert exported.status_code == 200
        assert exported.json()["case"]["code"] == "DATA-001"
        assert len(exported.json()["messages"]) == 2
        archived = client.patch(f"/cases/{case_id}", json={"archived": True})
        assert archived.json()["archived"] is True
        assert client.delete(f"/cases/{case_id}").status_code == 204
        assert client.get(f"/cases/{case_id}/stats").status_code == 404


def test_case_projects_keep_files_and_memory_isolated(monkeypatch):
    opened = []
    monkeypatch.setattr("app.main.open_in_file_manager", opened.append)
    with TestClient(app) as client:
        first = client.post("/cases", json={"name": "甲案", "code": "PROJECT-A"}).json()
        second = client.post("/cases", json={"name": "乙案", "code": "PROJECT-B"}).json()
        for case in (first, second):
            client.post(
                f"/cases/{case['id']}/sources",
                files={"file": (f"{case['code']}.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        first_conversation = client.post(
            f"/cases/{first['id']}/conversations", json={"title": "甲案研判"}
        ).json()
        second_conversation = client.post(
            f"/cases/{second['id']}/conversations", json={"title": "乙案研判"}
        ).json()
        client.patch(f"/conversations/{first_conversation['id']}", json={"analysis_enabled": True})
        client.post(
            f"/cases/{first['id']}/analysis",
            json={"question": "wxid_l3q7j5v9的登录地址是什么", "conversation_id": first_conversation["id"]},
        )

        first_root = settings.data_dir / first["id"]
        second_root = settings.data_dir / second["id"]
        assert first_root != second_root
        assert len(list((first_root / "data").glob("*.xlsx"))) == 1
        assert len(list((second_root / "data").glob("*.xlsx"))) == 1
        first_memory = json.loads((first_root / "conversations" / f"{first_conversation['id']}.json").read_text(encoding="utf-8"))
        second_memory = json.loads((second_root / "conversations" / f"{second_conversation['id']}.json").read_text(encoding="utf-8"))
        assert first_memory["case_code"] == "PROJECT-A"
        assert [turn["question"] for turn in first_memory["turns"]] == ["wxid_l3q7j5v9的登录地址是什么"]
        assert second_memory["case_code"] == "PROJECT-B"
        assert second_memory["turns"] == []

        updated = client.patch(
            f"/cases/{first['id']}", json={"pinned": True, "section": "研判中"}
        ).json()
        assert updated["pinned"] is True
        project = json.loads((first_root / "project.json").read_text(encoding="utf-8"))
        assert project["section"] == "研判中"
        folder = client.get(f"/cases/{first['id']}/folder").json()
        assert {item["path"] for item in folder["files"]} >= {
            "project.json",
            f"conversations/{first_conversation['id']}.json",
        }
        assert client.post(f"/cases/{first['id']}/folder/open").status_code == 200
        assert opened == [first_root]
        archived = client.post(f"/cases/{first['id']}/conversations/archive").json()
        assert archived["archived"] == 1

        assert client.delete(f"/cases/{first['id']}").status_code == 204
        assert not first_root.exists()
        assert second_root.exists()


def test_project_runtime_settings_skills_models_and_usage(monkeypatch):
    with TestClient(app) as client:
        case = client.post("/cases", json={"name": "配置测试", "code": "SET-001"}).json()
        case_id = case["id"]
        skill_dir = settings.data_dir / case_id / "agent" / "skills" / "case-helper"
        skill_dir.mkdir(parents=True)
        (skill_dir / "SKILL.md").write_text("---\nname: case-helper\n---\n只使用案件内材料。", encoding="utf-8")

        skills = client.get(f"/cases/{case_id}/skills").json()
        project_skill = next(item for item in skills if item["id"] == "项目:case-helper")
        assert project_skill["scope"] == "项目"
        assert all(item["enabled"] for item in skills)

        payload = client.get(f"/cases/{case_id}/settings").json()
        payload["skills"] = [project_skill["id"]]
        payload["system_prompts"] = [
            {"id": "builtin", "name": "案件专用研判规则", "enabled": False, "content": "", "builtin": True},
            {"id": "custom-1", "name": "案件提示词", "enabled": True, "content": "测试提示词", "builtin": False},
        ]
        payload["model"]["mode"] = "api"
        payload["model"]["api"] = {"base_url": "https://example.test/v1", "model": "model-a", "api_key": "secret"}
        payload["model"]["api_presets"] = [{
            "id": "preset-1",
            "name": "办公接口",
            "note": "自定义说明",
            "base_url": "https://example.test/v1",
            "model": "model-a",
            "api_key": "preset-secret",
        }]
        saved = client.put(f"/cases/{case_id}/settings", json=payload).json()
        assert saved["model"]["api"]["api_key_set"] is True
        assert "api_key" not in saved["model"]["api"]
        assert saved["model"]["api_presets"][0]["name"] == "办公接口"
        assert saved["model"]["api_presets"][0]["note"] == "自定义说明"
        assert saved["model"]["api_presets"][0]["api_key_set"] is True
        assert "api_key" not in saved["model"]["api_presets"][0]
        stored = json.loads((settings.data_dir / case_id / "agent" / "settings.json").read_text(encoding="utf-8"))
        assert stored["model"]["api"]["api_key"] == "secret"
        assert stored["model"]["api_presets"][0]["api_key"] == "preset-secret"
        assert client.get(f"/cases/{case_id}/settings").json()["skills"] == ["项目:case-helper"]
        assert client.get(f"/cases/{case_id}/settings").json()["system_prompts"][0]["name"] == "案件专用研判规则"
        assert client.get(f"/cases/{case_id}/settings").json()["system_prompts"][1]["enabled"] is True

        class ModelResponse:
            def raise_for_status(self):
                return None

            def json(self):
                return {"data": [{"id": "model-b"}, {"id": "model-a"}]}

        monkeypatch.setattr("app.runtime_settings.httpx.get", lambda *args, **kwargs: ModelResponse())
        models = client.post(f"/cases/{case_id}/models", json={"base_url": "https://example.test/v1", "api_key": "secret"}).json()
        assert models == {"models": ["model-a", "model-b"]}
        assert client.get(f"/cases/{case_id}/usage").json()["total_tokens"] == 0


def person_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员信息"
    sheet.append(["序号", "姓名", "微信昵称", "微信号", "性别", "身份证号", "手机号", "与老K关系", "快递地址", "外卖地址", "常用IP"])
    sheet.append([1, "周凯", "老K", "wxid_zhou", "男", "110101199001011234", "13800000000", "本人", "A路1号", "B路2号", "1.1.1.1"])
    sheet.append([2, "刘桂香", "香香", "wxid_liu", "女", "110101196001011234", "13900000000", "母亲", "A路1号", None, "1.1.1.2"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_import_persons_into_sqlite():
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "人员案", "code": "PERSON-001"}).json()["id"]
        imported = client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("info.xlsx", person_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert imported.status_code == 201
        assert imported.json()["imported_messages"] == 0
        assert imported.json()["imported_persons"] == 2
        graph = client.get(f"/cases/{case_id}/graph").json()
        assert {node["name"] for node in graph["nodes"]} >= {"周凯", "刘桂香"}
        assert client.delete(f"/cases/{case_id}").status_code == 204
        assert client.get("/cases").json() == []


def test_standard_retrieval_questions():
    questions = json.loads((Path(__file__).parent / "eval_questions.json").read_text(encoding="utf-8"))
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "检索基准", "code": "EVAL-001"}).json()["id"]
        client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("chat.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        for item in questions:
            hit = client.get(f"/cases/{case_id}/search", params={"q": item["question"], "limit": 1}).json()[0]
            for field, expected in item.items():
                if field.startswith("expected_"):
                    assert hit[field.removeprefix("expected_")] == expected


def test_conversation_materials_are_isolated():
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "材料范围", "code": "SCOPE-001"}).json()["id"]
        client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("case.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        first = client.post(f"/cases/{case_id}/conversations", json={"title": "对话甲"}).json()
        second = client.post(f"/cases/{case_id}/conversations", json={"title": "对话乙"}).json()
        assert first["analysis_enabled"] is False
        assert second["analysis_enabled"] is False
        for item in (first, second):
            enabled = client.patch(f"/conversations/{item['id']}", json={"analysis_enabled": True})
            assert enabled.json()["analysis_enabled"] is True
        scoped = client.post(
            f"/cases/{case_id}/sources",
            params={"conversation_id": first["id"]},
            files={"file": ("private.xlsx", conversation_only_workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert scoped.status_code == 201
        assert scoped.json()["scope"] == "conversation"
        assert scoped.json()["imported_messages"] == 1
        owned = client.post(
            f"/cases/{case_id}/analysis",
            json={"question": "对话专属句子", "conversation_id": first["id"], "limit": 10},
        )
        hidden = client.post(
            f"/cases/{case_id}/analysis",
            json={"question": "对话专属句子", "conversation_id": second["id"], "limit": 10},
        )
        assert "对话专属句子" in [item["content"] for item in owned.json()["evidence"]]
        assert "对话专属句子" not in [item["content"] for item in hidden.json()["evidence"]]
        shared = client.post(
            f"/cases/{case_id}/analysis",
            json={"question": "蓝盒还在原处吗", "conversation_id": second["id"], "limit": 10},
        )
        assert "蓝盒还在原处吗" in [item["content"] for item in shared.json()["evidence"]]
        groups = client.get(f"/cases/{case_id}/groups", params={"conversation_id": second["id"]}).json()
        assert "g-private" not in {item["group_id"] for item in groups}


def test_delete_conversation_keeps_case_materials_and_other_cases():
    payload = workbook_bytes()
    sheet = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with TestClient(app) as client:
        case_a = client.post("/cases", json={"name": "甲案", "code": "KEEP-A"}).json()
        case_b = client.post("/cases", json={"name": "乙案", "code": "KEEP-B"}).json()
        assert client.post(
            f"/cases/{case_a['id']}/sources",
            files={"file": ("chat.xlsx", payload, sheet)},
        ).status_code == 201
        assert client.post(
            f"/cases/{case_b['id']}/sources",
            files={"file": ("chat.xlsx", payload, sheet)},
        ).status_code == 201
        conversation = client.post(
            f"/cases/{case_a['id']}/conversations", json={"title": "研判甲"}
        ).json()
        scoped = client.post(
            f"/cases/{case_a['id']}/sources",
            params={"conversation_id": conversation["id"]},
            files={"file": ("private.xlsx", conversation_only_workbook_bytes(), sheet)},
        )
        assert scoped.status_code == 201
        assert client.delete(f"/conversations/{conversation['id']}").status_code == 204
        assert client.get(f"/cases/{case_a['id']}/conversations").json() == []
        stats_a = client.get(f"/cases/{case_a['id']}/stats").json()
        assert stats_a["messages"] == 3
        assert stats_a["source_files"] == 2
        assert stats_a["case_source_files"] == 2
        sources = client.get(f"/cases/{case_a['id']}/sources").json()
        assert all(item["conversation_id"] is None for item in sources)
        groups_a = {item["group_id"] for item in client.get(f"/cases/{case_a['id']}/groups").json()}
        assert groups_a == {"g1", "g-private"}
        stats_b = client.get(f"/cases/{case_b['id']}/stats").json()
        assert stats_b["messages"] == 2
        assert client.delete(f"/cases/{case_a['id']}").status_code == 204
        remaining = client.get("/cases").json()
        assert [item["code"] for item in remaining] == ["KEEP-B"]
        stats_b = client.get(f"/cases/{case_b['id']}/stats").json()
        assert stats_b["messages"] == 2
        groups_b = {item["group_id"] for item in client.get(f"/cases/{case_b['id']}/groups").json()}
        assert groups_b == {"g1"}


def test_chat_labels_and_persons():
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "标注案件", "code": "CASE-LABEL"}).json()["id"]
        empty = client.get(f"/cases/{case_id}/chat-labels")
        assert empty.status_code == 200
        assert empty.json()["groups"] == {}
        saved = client.put(
            f"/cases/{case_id}/chat-labels",
            json={
                "groups": {"g1": {"type": "同事群/工作群", "type_source": "manual", "hidden": True}},
                "senders": {"wxid_l3q7j5v9": {"name": "周凯", "name_source": "manual"}},
            },
        )
        assert saved.status_code == 200
        assert saved.json()["groups"]["g1"]["type"] == "同事群/工作群"
        assert saved.json()["groups"]["g1"]["hidden"] is True
        assert saved.json()["senders"]["wxid_l3q7j5v9"]["name"] == "周凯"
        persons = client.get(f"/cases/{case_id}/persons")
        assert persons.status_code == 200
        groups = client.get(f"/cases/{case_id}/groups")
        assert groups.status_code == 200


def friend_workbook_bytes(friends: list[tuple[str, str, str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "微信好友"
    sheet.append(["好友昵称", "备注", "微信ID", "微信号", "原始微信号"])
    for nickname, remark, wechat_id, alias in friends:
        sheet.append([nickname, remark, wechat_id, alias, ""])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def person_multi_wechat_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员信息"
    sheet.append(["序号", "姓名", "微信昵称", "微信ID", "微信号", "原始微信号", "性别", "身份证号", "手机号", "快递地址", "外卖地址"])
    sheet.append([1, "周凯", "老K", "wxid_zhou_a", "zhoukai01", "old_zhou", "男", "110101199001011234", "13800000000", "A路1号", "外卖路9号"])
    sheet.append([2, "周凯", "凯子", "wxid_zhou_b", "zhoukai02", "", "男", "110101199001011234", "13800000000", "A路1号", "外卖路9号"])
    sheet.append([3, "张三", "阿三", "wxid_zhang", "zhangsan01", "", "男", "110101199201011234", "13900000000", "C路3号", ""])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_wechat_friend_screen_and_graph_profile():
    sheet = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "好友案", "code": "FRIEND-001"}).json()["id"]
        persons = client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("人员信息.xlsx", person_multi_wechat_bytes(), sheet)},
        )
        assert persons.status_code == 201
        assert persons.json()["imported_persons"] == 3
        zhou = client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("周凯微信好友.xlsx", friend_workbook_bytes([
                ("阿三", "张三", "wxid_zhang", "zhangsan01"),
                ("李四", "同事", "wxid_li", "lisi01"),
                ("王五", "共同", "wxid_wang", "wangwu01"),
            ]), sheet)},
        )
        assert zhou.status_code == 201
        assert zhou.json()["imported_friends"] == 3
        zhang = client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("张三微信好友.xlsx", friend_workbook_bytes([
                ("老K", "周凯", "wxid_zhou_b", "zhoukai02"),
                ("王五", "共同", "wxid_wang", "wangwu01"),
            ]), sheet)},
        )
        assert zhang.status_code == 201
        graph = client.get(f"/cases/{case_id}/graph").json()
        zhou_node = next(node for node in graph["nodes"] if node["name"] == "周凯")
        assert zhou_node["takeout_address"] == "外卖路9号"
        assert zhou_node["delivery_address"] == "A路1号"
        assert {account["wechat_id"] for account in zhou_node["accounts"]} >= {"wxid_zhou_a", "wxid_zhou_b"}
        friend_edges = [edge for edge in graph["edges"] if edge.get("kind") == "friend"]
        assert friend_edges
        assert any(edge["mutual"] for edge in friend_edges)
        screened = client.post(f"/cases/{case_id}/friend-screen", json={"names": ["周凯", "张三"]})
        assert screened.status_code == 200
        body = screened.json()
        assert body["pairs"][0]["mutual"] is True
        assert any(item["wechat_id"] == "wxid_wang" for item in body["common_friends"])


def test_graph_person_and_relation_edits():
    sheet = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "编辑图谱", "code": "GRAPH-EDIT-001"}).json()["id"]
        imported = client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("人员信息.xlsx", person_multi_wechat_bytes(), sheet)},
        )
        assert imported.status_code == 201
        graph = client.get(f"/cases/{case_id}/graph").json()
        zhou = next(node for node in graph["nodes"] if node["name"] == "周凯")
        assert zhou["person_id"]
        updated = client.patch(
            f"/cases/{case_id}/persons/{zhou['person_id']}",
            json={"phone": "13900001111", "gender": "男"},
        )
        assert updated.status_code == 200
        created = client.post(
            f"/cases/{case_id}/persons",
            json={"name": "李四", "owner_name": "周凯", "relation": "同事", "phone": "13700000000"},
        )
        assert created.status_code == 201
        graph = client.get(f"/cases/{case_id}/graph").json()
        li = next(node for node in graph["nodes"] if node["name"] == "李四")
        assert li["phone"] == "13700000000"
        assert any(edge["relation"] == "同事" and li["id"] in {edge["source"], edge["target"]} for edge in graph["edges"])
        renamed = client.put(
            f"/cases/{case_id}/relations",
            json={"source": "周凯", "target": "李四", "relation": "朋友"},
        )
        assert renamed.status_code == 200
        graph = client.get(f"/cases/{case_id}/graph").json()
        assert any(edge["relation"] == "朋友" for edge in graph["edges"])
        deleted = client.delete(f"/cases/{case_id}/persons/{li['person_id']}")
        assert deleted.status_code == 204
        graph = client.get(f"/cases/{case_id}/graph").json()
        assert all(node["name"] != "李四" for node in graph["nodes"])
        related = client.post(
            f"/cases/{case_id}/persons",
            json={"name": "王敏", "owner_name": "周凯", "relation": "同事"},
        )
        assert related.status_code == 201
        wang = next(node for node in client.get(f"/cases/{case_id}/graph").json()["nodes"] if node["name"] == "王敏")
        assert wang["description"] != "本人"
        promoted = client.patch(f"/cases/{case_id}/persons/{wang['person_id']}", json={"as_subject": True})
        assert promoted.status_code == 200
        wang = next(node for node in client.get(f"/cases/{case_id}/graph").json()["nodes"] if node["name"] == "王敏")
        assert wang["description"] == "本人"
        restored = client.patch(f"/cases/{case_id}/persons/{wang['person_id']}", json={"remark": "周凯的同事"})
        assert restored.status_code == 200
        wang = next(node for node in client.get(f"/cases/{case_id}/graph").json()["nodes"] if node["name"] == "王敏")
        assert wang["description"] == "周凯的同事"
        removed_rel = client.put(
            f"/cases/{case_id}/relations",
            json={"source": "周凯", "target": "王敏", "relation": "同事", "removed": True},
        )
        assert removed_rel.status_code == 200
        graph = client.get(f"/cases/{case_id}/graph").json()
        zhou = next(node for node in graph["nodes"] if node["name"] == "周凯")
        wang = next(node for node in graph["nodes"] if node["name"] == "王敏")
        assert all({edge["source"], edge["target"]} != {zhou["id"], wang["id"]} for edge in graph["edges"])


def test_graph_keep_existing_edges_and_find_new_subject():
    sheet = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "连线保留", "code": "GRAPH-KEEP-001"}).json()["id"]
        client.post(
            f"/cases/{case_id}/sources",
            files={"file": ("人员信息.xlsx", person_multi_wechat_bytes(), sheet)},
        )
        mother = client.post(
            f"/cases/{case_id}/persons",
            json={"name": "刘桂香", "owner_name": "周凯", "relation": "母亲"},
        )
        assert mother.status_code == 201
        before = client.get(f"/cases/{case_id}/graph").json()
        zhou = next(node for node in before["nodes"] if node["name"] == "周凯")
        liu = next(node for node in before["nodes"] if node["name"] == "刘桂香")
        zhang = next(node for node in before["nodes"] if node["name"] == "张三")
        mother_edges = [
            edge for edge in before["edges"]
            if {edge["source"], edge["target"]} == {zhou["id"], liu["id"]}
        ]
        assert mother_edges
        linked = client.put(
            f"/cases/{case_id}/relations",
            json={"source": "张三", "target": "刘桂香", "relation": "同事"},
        )
        assert linked.status_code == 200
        after = client.get(f"/cases/{case_id}/graph").json()
        zhou = next(node for node in after["nodes"] if node["name"] == "周凯")
        liu = next(node for node in after["nodes"] if node["name"] == "刘桂香")
        zhang = next(node for node in after["nodes"] if node["name"] == "张三")
        assert liu["description"] == "周凯的母亲"
        assert any({edge["source"], edge["target"]} == {zhou["id"], liu["id"]} for edge in after["edges"])
        assert any({edge["source"], edge["target"]} == {zhang["id"], liu["id"]} and edge["relation"] == "同事" for edge in after["edges"])
        created = client.post(f"/cases/{case_id}/persons", json={"name": "测试当事人"})
        assert created.status_code == 201
        graph = client.get(f"/cases/{case_id}/graph").json()
        subject = next(node for node in graph["nodes"] if node["name"] == "测试当事人")
        assert subject["description"] == "本人"
        assert subject["person_id"] == created.json()["id"]
        promoted = client.post(f"/cases/{case_id}/persons", json={"name": "张三"})
        assert promoted.status_code == 201
        zhang = next(node for node in client.get(f"/cases/{case_id}/graph").json()["nodes"] if node["name"] == "张三")
        assert zhang["description"] == "本人"


def subject_relation_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员关系"
    sheet.append(["主体", "关系人", "关系类型"])
    sheet.append(["", "王五", "同事"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_graph_subject_create_delete_and_owned_import():
    sheet = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    with TestClient(app) as client:
        case_id = client.post("/cases", json={"name": "当事人栏", "code": "SUBJECT-BAR-001"}).json()["id"]
        created = client.post(f"/cases/{case_id}/persons", json={"name": "周凯", "phone": "13800001111"})
        assert created.status_code == 201
        assert created.json()["role"] == "subject"
        graph = client.get(f"/cases/{case_id}/graph").json()
        zhou = next(node for node in graph["nodes"] if node["name"] == "周凯")
        assert zhou["id"] in graph["subjects"]
        related = client.post(
            f"/cases/{case_id}/persons",
            json={"name": "刘桂香", "owner_name": "周凯", "relation": "母亲"},
        )
        assert related.status_code == 201
        blocked = client.delete(f"/cases/{case_id}/persons/{zhou['person_id']}")
        assert blocked.status_code == 422
        friends = client.post(
            f"/cases/{case_id}/sources?owner_name=周凯",
            files={"file": ("好友表.xlsx", friend_workbook_bytes([
                ("阿三", "张三", "wxid_zhang", "zhangsan01"),
            ]), sheet)},
        )
        assert friends.status_code == 201
        assert friends.json()["imported_friends"] == 1
        relations = client.post(
            f"/cases/{case_id}/sources?owner_name=周凯",
            files={"file": ("关系表.xlsx", subject_relation_bytes(), sheet)},
        )
        assert relations.status_code == 201
        assert relations.json()["imported_persons"] >= 1
        graph = client.get(f"/cases/{case_id}/graph").json()
        assert any(node["name"] == "王五" for node in graph["nodes"])
        for node in graph["nodes"]:
            if node["name"] == "周凯" and node.get("person_id"):
                removed = client.delete(f"/cases/{case_id}/persons/{node['person_id']}?as_subject=true")
                assert removed.status_code == 204
        graph = client.get(f"/cases/{case_id}/graph").json()
        assert all(node["name"] not in {"周凯", "刘桂香", "王五"} for node in graph["nodes"])

